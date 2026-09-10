"""File reading tool — PDF, Word, Excel, plain text."""
import pathlib
from ._base import JarvisTool, ToolResult


class FilesTool(JarvisTool):
    name = "read_file"
    description = "Read a file from disk — supports PDF, Word, Excel, and plain text"
    tags = ["file", "read", "pdf", "document", "excel", "word", "txt", "open", "load"]

    def run(self, path: str = "", max_chars: int = 8000, **kwargs) -> ToolResult:
        if not path:
            return ToolResult(False, "", "path is required", self.name)

        p = pathlib.Path(path)
        if not p.exists():
            return ToolResult(False, "", f"File not found: {path}", self.name)

        suffix = p.suffix.lower()
        try:
            if suffix == ".pdf":
                import pdfplumber
                parts = []
                with pdfplumber.open(path) as pdf:
                    for page in pdf.pages[:10]:
                        parts.append(page.extract_text() or "")
                return ToolResult(True, "\n".join(parts)[:max_chars], tool_name=self.name)

            if suffix in (".docx", ".doc"):
                from docx import Document
                doc = Document(path)
                text = "\n".join(p.text for p in doc.paragraphs)
                return ToolResult(True, text[:max_chars], tool_name=self.name)

            if suffix in (".xlsx", ".xls"):
                import openpyxl
                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                lines = []
                for sheet in wb.worksheets[:3]:
                    for row in sheet.iter_rows(max_row=200, values_only=True):
                        lines.append("\t".join(str(c) for c in row if c is not None))
                return ToolResult(True, "\n".join(lines)[:max_chars], tool_name=self.name)

            text = p.read_text(encoding="utf-8", errors="replace")
            return ToolResult(True, text[:max_chars], tool_name=self.name)

        except Exception as e:
            return ToolResult(False, "", str(e), self.name)
