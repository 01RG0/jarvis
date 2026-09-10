"""File search and document reading."""
from __future__ import annotations

import os
import pathlib

from . import register


def search_files(pattern: str, root: str = "C:/Users") -> list[str]:
    """Glob search. Falls back gracefully if root doesn't exist."""
    base = pathlib.Path(root)
    if not base.exists():
        base = pathlib.Path.home()
    try:
        return [str(p) for p in base.rglob(pattern)][:50]
    except PermissionError:
        return []


def read_file(path: str, max_chars: int = 8000) -> str:
    p = pathlib.Path(path)
    if not p.exists():
        return f"File not found: {path}"

    suffix = p.suffix.lower()
    try:
        if suffix == ".pdf":
            import pdfplumber
            text_parts = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages[:10]:
                    text_parts.append(page.extract_text() or "")
            return "\n".join(text_parts)[:max_chars]

        if suffix in (".docx", ".doc"):
            from docx import Document
            doc = Document(path)
            return "\n".join(p.text for p in doc.paragraphs)[:max_chars]

        if suffix in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            lines = []
            for sheet in wb.worksheets[:3]:
                for row in sheet.iter_rows(max_row=100, values_only=True):
                    lines.append("\t".join(str(c) for c in row if c is not None))
            return "\n".join(lines)[:max_chars]

        # Plain text fallback
        return p.read_text(encoding="utf-8", errors="replace")[:max_chars]

    except Exception as e:
        return f"Read error: {e}"


def write_file(path: str, content: str) -> None:
    pathlib.Path(path).write_text(content, encoding="utf-8")


def delete_file(path: str) -> None:
    os.remove(path)


register("search_files", search_files, tier="safe")
register("list_files",   search_files, tier="safe")  # alias
register("read_file",    read_file,    tier="medium")
register("write_file",   write_file,   tier="dangerous")
register("delete_file",  delete_file,  tier="dangerous")
